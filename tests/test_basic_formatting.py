import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import core.config_manager as config_manager
from format_conversion import convert_text_to_docx, reformat_docx


class BasicFormattingTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_config_path = config_manager._CONFIG_PATH
        config_manager._CONFIG_PATH = Path(self.temp_dir.name) / "config.json"

    def tearDown(self):
        config_manager._CONFIG_PATH = self.original_config_path
        self.temp_dir.cleanup()

    def test_default_heading_one_matches_required_format(self):
        settings = config_manager.get_heading_settings()["1"]
        self.assertEqual(settings["font_cn"], "黑体")
        self.assertEqual(settings["size_name"], "小三")
        self.assertEqual(settings["size_pt"], 15.0)
        self.assertEqual(settings["space_before_pt"], 40.0)
        self.assertEqual(settings["space_after_pt"], 20.0)
        self.assertTrue(settings["center"])

        scene = config_manager.get_active_scene_config()
        heading = scene.styles["heading1"]
        self.assertTrue(heading.page_break_before)
        self.assertTrue(heading.keep_with_next)

    def test_heading_overrides_are_saved_without_legacy_modes(self):
        settings = config_manager.get_heading_settings()
        settings["2"].update(
            font_cn="楷体",
            size_name="三号",
            space_before_pt=18,
            space_after_pt=9,
            center=True,
        )
        config_manager.save_heading_settings(settings)

        saved = config_manager.load_full_config()
        self.assertEqual(set(saved), {"heading_settings"})
        scene = config_manager.get_active_scene_config()
        heading = scene.styles["heading2"]
        self.assertEqual(heading.font_cn, "楷体")
        self.assertEqual(heading.size_pt, 16)
        self.assertEqual(heading.space_before_pt, 18)
        self.assertEqual(heading.space_after_pt, 9)
        self.assertEqual(heading.alignment, "center")

    def test_old_numeric_font_size_is_migrated_to_chinese_size_name(self):
        config_manager.save_full_config(
            {"heading_settings": {"1": {"font_cn": "黑体", "size_pt": 15}}}
        )
        settings = config_manager.get_heading_settings()["1"]
        self.assertEqual(settings["size_name"], "小三")
        self.assertEqual(settings["size_pt"], 15.0)

    def test_markdown_heading_one_is_applied_and_no_blank_cover_is_added(self):
        output = Path(self.temp_dir.name) / "heading.docx"
        convert_text_to_docx("# 第一章 绪论\n\n正文内容。", str(output))
        doc = Document(output)
        non_empty = [p for p in doc.paragraphs if p.text.strip()]
        self.assertTrue(non_empty)
        heading = non_empty[0]

        self.assertIn("第一章", heading.text)
        self.assertEqual(heading.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(heading.paragraph_format.space_before.pt, 40.0)
        self.assertEqual(heading.paragraph_format.space_after.pt, 20.0)
        self.assertTrue(heading.paragraph_format.page_break_before)
        self.assertTrue(heading.paragraph_format.keep_with_next)
        self.assertEqual(heading.runs[0].font.size.pt, 15.0)
        r_fonts = heading.runs[0]._r.get_or_add_rPr().find(qn("w:rFonts"))
        self.assertIsNotNone(r_fonts)
        self.assertEqual(r_fonts.get(qn("w:eastAsia")), "黑体")

    def test_markdown_toc_section_is_omitted(self):
        output = Path(self.temp_dir.name) / "without_toc.docx"
        text = "# 目录\n\n第一章 绪论……1\n\n# 第一章 绪论\n\n正文内容。"
        convert_text_to_docx(text, str(output))
        rendered = "\n".join(p.text for p in Document(output).paragraphs)
        self.assertNotIn("目录", rendered)
        self.assertNotIn("绪论……1", rendered)
        self.assertIn("第一章 绪论", rendered)
        self.assertIn("正文内容", rendered)

    def test_existing_docx_toc_section_is_omitted(self):
        source = Path(self.temp_dir.name) / "source.docx"
        output = Path(self.temp_dir.name) / "output.docx"
        doc = Document()
        doc.add_heading("目录", level=1)
        doc.add_paragraph("第一章 绪论........1")
        doc.add_heading("第一章 绪论", level=1)
        doc.add_paragraph("这是正文。")
        doc.save(source)

        reformat_docx(str(source), str(output))
        rendered = "\n".join(p.text for p in Document(output).paragraphs)
        self.assertNotIn("目录", rendered)
        self.assertNotIn("绪论........1", rendered)
        self.assertIn("第一章 绪论", rendered)
        self.assertIn("这是正文", rendered)


if __name__ == "__main__":
    unittest.main()
